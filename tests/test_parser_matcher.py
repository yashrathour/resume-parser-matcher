"""
Verification test script for Resume Parser and Skill Matcher modules.
"""

import io
import os
import sys

# Ensure project root is in python path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

import docx
from utils.parser import (
    extract_text_from_pdf,
    extract_text_from_docx,
    extract_text_from_file,
    extract_contact_info,
    extract_name,
    extract_skills,
    get_spacy_model,
    DEFAULT_SKILLS
)
from utils.matcher import (
    calculate_semantic_score,
    compare_skills
)

def test_docx_generation_and_extraction():
    print("--- Testing DOCX extraction ---")
    doc = docx.Document()
    doc.add_heading("Alex Morgan", level=0)
    doc.add_paragraph("alex.morgan@techcorp.io | +1 (415) 890-1234 | San Francisco, CA")
    doc.add_heading("Professional Summary", level=1)
    doc.add_paragraph("Senior Full-Stack Engineer with 6+ years of experience building scalable microservices in Python, Django, React, and PostgreSQL.")
    doc.add_heading("Skills", level=1)
    table = doc.add_table(rows=1, cols=2)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Languages & Cloud"
    hdr_cells[1].text = "Python, TypeScript, Go, Docker, Kubernetes, AWS"

    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)

    extracted_text = extract_text_from_docx(stream)
    assert "Alex Morgan" in extracted_text
    assert "alex.morgan@techcorp.io" in extracted_text
    assert "Kubernetes" in extracted_text
    print("✓ DOCX extraction passed.")
    return extracted_text

def test_pdf_extraction():
    print("\n--- Testing PDF extraction ---")
    minimal_pdf = b"""%PDF-1.4
1 0 obj
<< /Type /Catalog /Pages 2 0 R >>
endobj
2 0 obj
<< /Type /Pages /Kids [3 0 R] /Count 1 >>
endobj
3 0 obj
<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>
endobj
4 0 obj
<< /Length 55 >>
stream
BT
/F1 12 Tf
100 700 Td
(Jane Doe jane@example.com) Tj
ET
endstream
endobj
5 0 obj
<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>
endobj
xref
0 6
0000000000 65535 f 
0000000009 00000 n 
0000000058 00000 n 
0000000115 00000 n 
0000000224 00000 n 
0000000330 00000 n 
trailer
<< /Size 6 /Root 1 0 R >>
startxref
403
%%EOF"""

    text = extract_text_from_pdf(io.BytesIO(minimal_pdf))
    assert "Jane Doe" in text
    assert "jane@example.com" in text
    print("✓ PDF extraction passed.")

def test_information_extraction(sample_text):
    print("\n--- Testing Information Extraction ---")
    nlp = get_spacy_model()
    name = extract_name(sample_text, nlp=nlp)
    print(f"Extracted Name: {name}")
    assert name == "Alex Morgan", f"Expected 'Alex Morgan', got '{name}'"

    contact = extract_contact_info(sample_text)
    print(f"Extracted Contact: {contact}")
    assert contact["email"] == "alex.morgan@techcorp.io"
    assert "890-1234" in contact["phone"]

    skills = extract_skills(sample_text, DEFAULT_SKILLS)
    print(f"Extracted Skills: {skills}")
    for expected_skill in ["Python", "Django", "React", "PostgreSQL", "Docker", "Kubernetes", "AWS", "TypeScript", "Go"]:
        assert expected_skill in skills, f"Missing expected skill {expected_skill} in {skills}"
    
    # Check that 'C' was not falsely matched inside words
    assert "C" not in skills, "Skill 'C' was falsely matched!"
    print("✓ Information extraction passed.")
    return skills

def test_matching_logic(resume_skills):
    print("\n--- Testing Matching Logic ---")
    jd_text = """
    We are looking for a Senior Backend Engineer.
    Requirements:
    - 5+ years with Python, Django, Docker, and AWS.
    - Strong knowledge of PostgreSQL and Redis.
    - Hands-on experience with Terraform and CI/CD pipelines.
    - Frontend knowledge with React is a plus.
    """
    resume_text = """
    Alex Morgan
    alex.morgan@techcorp.io | +1 (415) 890-1234
    Experience with Python, Django, React, PostgreSQL, Docker, Kubernetes, AWS, Go, TypeScript.
    """

    jd_skills = extract_skills(jd_text, DEFAULT_SKILLS)
    print(f"JD Skills: {jd_skills}")

    # Semantic score
    score = calculate_semantic_score(resume_text, jd_text)
    print(f"Semantic match score: {score}%")
    assert 0.0 <= score <= 100.0
    assert score > 10.0, f"Expected reasonable score (> 10%), got {score}"

    # Skill comparison
    result = compare_skills(resume_skills, jd_skills)
    matched, missing = compare_skills(resume_skills, jd_skills)
    print(f"Matched Skills ({len(matched)}): {matched}")
    print(f"Missing Skills ({len(missing)}): {missing}")

    assert "Python" in matched
    assert "Django" in matched
    assert "Docker" in matched
    assert "AWS" in matched
    assert "PostgreSQL" in matched
    assert "React" in matched
    assert "Redis" in missing
    assert "Terraform" in missing
    assert "CI/CD" in missing

    # Test dictionary-like access
    assert result["Matched Skills"] == matched
    assert result["Missing Skills"] == missing
    print("✓ Matching logic passed.")

def test_edge_cases():
    print("\n--- Testing Edge Cases ---")
    # Empty inputs
    assert calculate_semantic_score("", "") == 0.0
    assert calculate_semantic_score("   ", "   ") == 0.0
    assert extract_contact_info("")["email"] == "Not found"
    assert extract_contact_info("")["phone"] == "Not found"
    assert extract_name("") == "Not found"
    assert extract_skills("") == []

    # Corrupted / invalid docx
    try:
        extract_text_from_docx(io.BytesIO(b"Not a valid zip/docx"))
        assert False, "Should have raised ValueError on corrupted docx"
    except ValueError as e:
        print(f"Caught expected error for corrupted docx: {e}")

    # Fallback when JD has no skills
    empty_jd_skills = []
    res = compare_skills(["Python", "SQL"], empty_jd_skills)
    assert res.matched_skills == []
    assert res.missing_skills == []
    print("✓ Edge cases passed.")

if __name__ == "__main__":
    extracted = test_docx_generation_and_extraction()
    test_pdf_extraction()
    skills = test_information_extraction(extracted)
    test_matching_logic(skills)
    test_edge_cases()
    print("\n🎉 ALL TESTS PASSED SUCCESSFULLY!")
